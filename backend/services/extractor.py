"""
Text extraction service for various file types
"""
import os
from typing import Optional


def extract_text(file_path: str, file_type: str) -> str:
    """
    Extract text content from various file types
    
    Args:
        file_path: Path to the file
        file_type: Type of file (txt, pdf, docx, etc.)
    
    Returns:
        Extracted text content
    """
    try:
        if file_type == 'txt' or file_type == 'md':
            return extract_from_txt(file_path)
        elif file_type == 'pdf':
            return extract_from_pdf(file_path)
        elif file_type in ['docx', 'doc']:
            return extract_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")
    except Exception as e:
        raise Exception(f"Error extracting text from {file_type}: {str(e)}")


def extract_from_txt(file_path: str) -> str:
    """Extract text from plain text or markdown files"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    except UnicodeDecodeError:
        # Try with different encoding
        with open(file_path, 'r', encoding='latin-1') as f:
            return f.read()


def extract_from_pdf(file_path: str) -> str:
    """Extract text from PDF files"""
    try:
        import PyPDF2
        
        text = []
        with open(file_path, 'rb') as f:
            pdf_reader = PyPDF2.PdfReader(f)
            num_pages = len(pdf_reader.pages)
            
            for page_num in range(num_pages):
                page = pdf_reader.pages[page_num]
                page_text = page.extract_text()
                if page_text:
                    text.append(page_text)
        
        return '\n\n'.join(text)
    except Exception as e:
        raise Exception(f"Error reading PDF: {str(e)}")


def extract_from_docx(file_path: str) -> str:
    """Extract text from DOCX files"""
    try:
        from docx import Document
        
        doc = Document(file_path)
        text = []
        
        # Extract paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)
        
        return '\n\n'.join(text)
    except Exception as e:
        raise Exception(f"Error reading DOCX: {str(e)}")


def extract_metadata(file_path: str, file_type: str) -> dict:
    """
    Extract metadata from files
    
    Returns:
        Dictionary containing metadata like page count, author, etc.
    """
    metadata = {
        'file_path': file_path,
        'file_type': file_type,
        'file_size': os.path.getsize(file_path)
    }
    
    try:
        if file_type == 'pdf':
            import PyPDF2
            with open(file_path, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)
                metadata['num_pages'] = len(pdf_reader.pages)
                if pdf_reader.metadata:
                    metadata['title'] = pdf_reader.metadata.get('/Title', '')
                    metadata['author'] = pdf_reader.metadata.get('/Author', '')
        
        elif file_type == 'docx':
            from docx import Document
            doc = Document(file_path)
            core_properties = doc.core_properties
            metadata['author'] = core_properties.author or ''
            metadata['title'] = core_properties.title or ''
            metadata['num_paragraphs'] = len(doc.paragraphs)
    
    except Exception as e:
        print(f"Warning: Could not extract metadata: {e}")
    
    return metadata
